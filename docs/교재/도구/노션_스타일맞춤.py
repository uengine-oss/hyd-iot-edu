"""기존 PPTX의 글자·캡처·배치는 그대로 두고 표지와 색·글꼴만 lecture-deck 디자인에 맞춘다.

사용: python 노션_스타일맞춤.py <원본.pptx> <출력.pptx> --course "피지컬 AI · 통합반 1일차" --subtitle "…"
원본은 읽기만 한다. 표지(첫 장)는 다른 강의의 것이므로 이 과정의 표지 틀로 바꾼다: 과정 표시와 부제는 인자로 받고,
큰 제목(강의 주제)만 원본을 그대로 둔다. 그 밖의 장은 저장한 뒤 원본과 대조해
글자(문단 단위)·그림(해시·위치·크기·자르기)이 하나라도 다르면 실패한다.
디자인 값(색·글꼴·표지 틀·넘기는 상자)은 lecture-deck의 deck2pptx.py에서 가져온다.
"""
from __future__ import annotations

import argparse
import hashlib
import sys
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

sys.path.insert(0, str(Path(__file__).resolve().parents[3] / ".claude" / "skills" / "lecture-deck" / "scripts"))
import deck2pptx as D  # noqa: E402

OLD_NAVY, OLD_ORANGE = RGBColor(0x12, 0x30, 0x5A), RGBColor(0xD4, 0x54, 0x1E)
# 원본 색 → 디자인 색. 표에 없는 색을 만나면 멈춘다(조용히 남기지 않는다).
TEXT_KEEP = {"FFFFFF", "B9C4D2"}
TEXT_MAP = {"3B4350": D.MUTED, "6B7280": D.MUTED, "000000": D.NAVY,
            "C3D2E6": D.COVER_SUB, "9DB4D4": D.SKY_ON_DARK}
BODY = {"111111", "222222"}            # 본문 글자: 굵기·문단 구성에 따라 INK / NAVY / SKY_D
LINE_MAP = {"E2E6EB": D.LINE, "C9D1DB": D.LINE, "D9DEE5": D.LINE, "DDE3EA": D.LINE}
FILL_MAP = {"12305A": D.NAVY, "F1F4F8": D.PANEL}
NEXT_Y, NEXT_H = 6.66, 0.58           # 넘기는 상자. 원본 캡처가 6.5in까지 내려와 기준(6.28)보다 아래에 둔다
COVER_ICONS = {"first": ("notebook", "pencil", "circle-check"), "last": ("circle-check", "pencil", "link")}


class RestyleError(Exception):
    pass


def rgb_of(color) -> str | None:
    return str(color.rgb) if color is not None and color.type is not None else None


def set_typeface(run):
    """글꼴 이름만 덱 글꼴로(코드 글꼴 Consolas는 둔다). 크기·굵기는 건드리지 않는다."""
    if run.font.name == "Consolas":
        return
    run.font.name = D.FONT
    rpr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        el = rpr.find(qn(tag))
        if el is None:
            el = rpr.makeelement(qn(tag), {})
            rpr.append(el)
        el.set("typeface", D.FONT)


def recolor_text(sh, where):
    for p in sh.text_frame.paragraphs:
        runs = [r for r in p.runs if r.text.strip()]
        mixed = any(r.font.bold for r in runs) and any(not r.font.bold for r in runs)
        for r in p.runs:
            set_typeface(r)
            c = rgb_of(r.font.color)
            if c is None or c in TEXT_KEEP:
                continue
            if c in BODY:
                r.font.color.rgb = (D.SKY_D if mixed else D.NAVY) if r.font.bold else D.INK
            elif c in TEXT_MAP:
                r.font.color.rgb = TEXT_MAP[c]
            else:
                raise RestyleError(f"{where} 모르는 글자색 {c}: {r.text[:20]}")


def recolor_shape(sh, where):
    if sh.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE and sh.fill.type == 1:
        c = str(sh.fill.fore_color.rgb)
        if c in FILL_MAP:
            sh.fill.fore_color.rgb = FILL_MAP[c]
        elif c != str(OLD_ORANGE):     # 주황은 캡처 위 표시(화면의 일부)라 그대로 둔다
            raise RestyleError(f"{where} 모르는 채움색 {c}")
    if sh.shape_type in (MSO_SHAPE_TYPE.AUTO_SHAPE, MSO_SHAPE_TYPE.LINE) and sh.line.fill.type == 1:
        c = str(sh.line.color.rgb)
        if c in LINE_MAP:
            sh.line.color.rgb = LINE_MAP[c]
        elif c != str(OLD_ORANGE):
            raise RestyleError(f"{where} 모르는 선색 {c}")


def place(sh, x, y, w, h):
    sh.left, sh.top, sh.width, sh.height = Inches(x), Inches(y), Inches(w), Inches(h)


def set_one_line(sh, text):
    """글상자의 글을 한 문단 한 줄로 바꾼다(첫 run의 서식은 유지)."""
    paras = sh.text_frame.paragraphs
    for p in paras[1:]:
        p._p.getparent().remove(p._p)
    runs = paras[0].runs
    for r in runs[1:]:
        r._r.getparent().remove(r._r)
    runs[0].text = text


def restyle_cover(sl, where, icons, cover=None):
    """남색 장(표지·마무리): 자리와 크기를 deck2pptx 표지 틀로. 글자는 그대로 두되,
    `cover`(과정 표시·부제)가 오면 표지 틀의 세 칸(과정·큰 제목·부제)만 남긴다 — 큰 제목은 원본 그대로."""
    D.set_bg(sl, D.NAVY)
    boxes = [sh for sh in sl.shapes if sh.has_text_frame and sh.text_frame.text.strip()]
    bars = [sh for sh in sl.shapes if sh.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE and sh.fill.type == 1
            and sh.fill.fore_color.rgb == OLD_ORANGE]
    if len(boxes) != 4 or len(bars) != 1:
        raise RestyleError(f"{where} 표지 구성이 예상과 다름(글상자 {len(boxes)}, 띠 {len(bars)})")
    label, title, sub, foot = boxes
    if cover:
        for key, limit in (("course", "제목"), ("subtitle", "부제")):
            if len(cover[key]) > D.LIMITS[limit if key == "subtitle" else "부제"]:
                raise RestyleError(f"{where} 표지 {key} 글자 수 초과: {cover[key]}")
        if "피지컬 AI" not in cover["course"]:
            raise RestyleError(f"{where} 표지 과정 표시에 「피지컬 AI」가 없음")
        set_one_line(label, cover["course"])
        set_one_line(sub, cover["subtitle"])
        foot._element.getparent().remove(foot._element)
    bars[0].fill.fore_color.rgb = D.SKY
    place(bars[0], D.MX, 1.55, 0.9, 0.07)
    rows = [(label, 1.75, 0.5, 18), (title, 2.45, 2.3, 40), (sub, 5.0 if cover else 4.9, 0.9 if cover else 0.5, 19)]
    if not cover:
        rows.append((foot, 5.6, 0.9, None))
    for sh, y, h, size in rows:
        place(sh, D.MX - 0.1, y, 8.8, h)
        recolor_text(sh, where)
        if size:
            for p in sh.text_frame.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(size)
    t, gap, cx = 1.3, 0.5, 11.0     # deck2pptx.render_cover의 아이콘 세로 줄과 같은 값
    top = (D.SLIDE_H - (t * len(icons) + gap * (len(icons) - 1))) / 2
    for i, name in enumerate(icons):
        y, last = top + i * (t + gap), i == len(icons) - 1
        D.add_box(sl, cx - t / 2, y, t, t, D.SKY if last else D.NAVY_2, shape=MSO_SHAPE.RECTANGLE)
        D.add_icon(sl, name, cx, y + t / 2, t * 0.6, D.NAVY if last else D.SKY, 1.5)
        if not last:
            D.add_box(sl, cx - 0.17, y + t + 0.08, 0.34, gap - 0.16, D.SKY, shape=MSO_SHAPE.DOWN_ARROW)


def restyle_body(sl, where):
    shapes = list(sl.shapes)
    title = next((sh for sh in shapes if sh.has_text_frame and sh.text_frame.text.strip()), None)
    if title is None or Emu(title.top).inches > 1.0:
        raise RestyleError(f"{where} 맨 위 제목 글상자를 찾지 못함")
    for sh in shapes:
        recolor_shape(sh, where)
        if sh.has_text_frame and sh.text_frame.text.strip():
            recolor_text(sh, where)
    paras = title.text_frame.paragraphs
    for r in paras[0].runs:            # 제목: 디자인 기준 크기·남색
        r.font.size, r.font.color.rgb = Pt(D.TITLE_PT), D.NAVY
    for p in paras[1:]:                # 제목 아래 설명 줄
        for r in p.runs:
            r.font.size, r.font.color.rgb = Pt(17), D.MUTED
    # 맨 아래 한 줄 → 넘기는 상자(테두리 상자 안 가운데 굵은 글씨). 문구는 그대로.
    foot = [sh for sh in shapes if sh.has_text_frame and sh.text_frame.text.strip() and Emu(sh.top).inches > 6.6]
    rule = [sh for sh in shapes if sh.shape_type == MSO_SHAPE_TYPE.LINE and abs(Emu(sh.top).inches - 6.5) < 0.02]
    if len(foot) != 1 or len(rule) != 1:
        raise RestyleError(f"{where} 아래 한 줄 구성이 예상과 다름(글상자 {len(foot)}, 구분선 {len(rule)})")
    foot, rule = foot[0], rule[0]
    x, w = Emu(rule.left).inches, Emu(rule.width).inches
    rule._element.getparent().remove(rule._element)
    box = D.add_box(sl, x, NEXT_Y, w, NEXT_H, D.WHITE, line=(D.NAVY, 1.75), shape=MSO_SHAPE.RECTANGLE)
    foot._element.addprevious(box._element)      # 상자를 글자 뒤에 둔다
    place(foot, x + 0.2, NEXT_Y, w - 0.4, NEXT_H)
    tf = foot.text_frame
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    for p in tf.paragraphs:
        p.alignment = PP_ALIGN.CENTER
        for r in p.runs:
            r.font.size, r.font.bold, r.font.color.rgb = Pt(18), True, D.NAVY


def snapshot(path: Path):
    """보존 대조용: 장마다 문단 글자 목록과 그림(해시·위치·크기·자르기) 목록."""
    out = []
    for sl in Presentation(path).slides:
        texts = sorted(p.text for sh in sl.shapes if sh.has_text_frame for p in sh.text_frame.paragraphs if p.text.strip())
        pics = sorted((hashlib.sha1(sh.image.blob).hexdigest(), sh.left, sh.top, sh.width, sh.height,
                       sh.crop_left, sh.crop_top, sh.crop_right, sh.crop_bottom)
                      for sh in sl.shapes if sh.shape_type == MSO_SHAPE_TYPE.PICTURE and not sh.name.startswith("아이콘:"))
        marks = sorted((sh.left, sh.top, sh.width, sh.height) for sh in sl.shapes
                       if sh.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE
                       and (sh.fill.type == 1 and sh.fill.fore_color.rgb == OLD_ORANGE
                            or sh.line.fill.type == 1 and sh.line.color.rgb == OLD_ORANGE))
        out.append((texts, pics, marks))
    return out


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("src")
    ap.add_argument("out")
    ap.add_argument("--course", required=True, help="표지 과정 표시(「피지컬 AI」 포함)")
    ap.add_argument("--subtitle", required=True, help="표지 부제 한 줄")
    a = ap.parse_args()
    src, out = Path(a.src), Path(a.out)
    cover = {"course": a.course, "subtitle": a.subtitle}
    if out.resolve() == src.resolve():
        print("원본과 출력이 같음", file=sys.stderr)
        return 2
    prs = Presentation(src)
    last = len(prs.slides)
    try:
        for i, sl in enumerate(prs.slides, 1):
            where = f"SLIDE {i:02d}"
            bg = sl.background.fill
            dark = bg.type == 1 and bg.fore_color.rgb == OLD_NAVY
            if dark and i in (1, last):
                restyle_cover(sl, where, COVER_ICONS["first" if i == 1 else "last"], cover if i == 1 else None)
            elif dark:
                raise RestyleError(f"{where} 표지·마무리가 아닌 남색 장")
            else:
                restyle_body(sl, where)
        prs.save(out)
        before, after = snapshot(src), snapshot(out)
        diffs = []
        for i, (b, a) in enumerate(zip(before, after), 1):
            b_marks = b[2] if i not in (1, last) else []      # 표지·마무리의 주황 띠는 하늘색 띠로 바꿨다
            if i == 1:      # 표지는 이 과정의 것으로 바꿨다. 큰 제목(강의 주제)만 원본과 같아야 한다
                title = [p.text for p in list(prs.slides)[0].shapes[1].text_frame.paragraphs]
                if not set(title) <= set(b[0]):
                    diffs.append("SLIDE 01 표지 큰 제목이 원본과 다름")
                continue
            for name, x, y in (("글자", b[0], a[0]), ("그림", b[1], a[1]), ("캡처 위 표시", b_marks, a[2])):
                if x != y:
                    diffs.append(f"SLIDE {i:02d} {name}가 원본과 다름")
        if len(before) != len(after) or diffs:
            raise RestyleError("보존 대조 실패:\n  " + "\n  ".join(diffs or ["장 수가 다름"]))
    except RestyleError as e:
        print(f"[실패] {e}", file=sys.stderr)
        return 1
    n_pic = sum(len(s[1]) for s in after)
    n_txt = sum(len(s[0]) for s in after)
    print(f"[완료] {out} — {last}장. 표지는 이 과정의 틀로 교체, 2~{last}장은 문단·그림 {n_pic}개·캡처 위 표시가 원본과 같음")
    return 0


if __name__ == "__main__":
    sys.exit(main())
