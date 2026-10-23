"""대본 → 발표자 노트. 대본 구획의 낭독 문단을 공백 하나로 이어 한 문단으로 만든다(`>` 줄 · `---` · 표 줄 · 굵게 표시 제외).
`SLIDE NN~MM` 묶음은 첫 장에 노트를 두고 뒤 장은 첫 장을 가리킨다(기준_이론대본 87).

    python3 notes.py check <대본 폴더> <PPTX> [<PPTX> …]     # 지금 노트가 대본과 같은 장 · 다른 장을 센다
    python3 notes.py write <대본 폴더> <PPTX 입력> <PPTX 출력>   # 다른 장의 노트를 대본대로 바꾼다(화면은 건드리지 않음)
"""
import re
import sys
from pathlib import Path
from pptx import Presentation

HEAD = re.compile(r"^## \[SLIDE ([0-9~]+)\].*$", re.M)


def expected(folder):
    """{장 번호: 노트 글}"""
    out = {}
    for f in sorted(Path(folder).glob("*.md")):
        t = f.read_text("utf-8")
        for m in HEAD.finditer(t):
            nxt = re.compile(r"^#{1,2} ", re.M).search(t, m.end())
            body = t[m.end():nxt.start() if nxt else len(t)]
            paras = []
            for p in re.split(r"\n\s*\n", body):
                lines = [x for x in p.split("\n") if x.strip() and not re.match(r"^(>|---\s*$|\|)", x)]
                if lines: paras.append(" ".join(x.strip() for x in lines).replace("**", ""))
            parts = m.group(1).split("~"); first, last = int(parts[0]), int(parts[-1])
            out[first] = " ".join(paras)
            for k in range(first + 1, last + 1):
                out[k] = f"{first}~{last}장 통합 발화는 {first}장 발표자 노트를 따른다."
    return out


def current(pptx):
    prs = Presentation(pptx)
    return prs, {i: (s.notes_slide.notes_text_frame.text if s.has_notes_slide else "") for i, s in enumerate(prs.slides, 1)}


if __name__ == "__main__":
    mode, folder = sys.argv[1], sys.argv[2]
    exp = expected(folder)
    if mode == "check":
        for pptx in sys.argv[3:]:
            _, cur = current(pptx)
            same = [n for n in exp if cur.get(n, "").strip() == exp[n].strip()]
            diff = [n for n in exp if n in cur and cur[n].strip() != exp[n].strip()]
            print(f"{Path(pptx).name}: 대본과 같은 노트 {len(same)}장, 다른 노트 {len(diff)}장 {diff}")
    elif mode == "write":
        prs, cur = current(sys.argv[3])
        changed = []
        for n, text in exp.items():
            s = prs.slides[n - 1]
            tf = s.notes_slide.notes_text_frame
            if tf.text.strip() == text.strip(): continue
            p0 = tf.paragraphs[0]
            for extra in tf.paragraphs[1:]: extra._p.getparent().remove(extra._p)   # 문단 하나로
            runs = p0.runs
            if runs:
                runs[0].text = text
                for r in runs[1:]: r._r.getparent().remove(r._r)
            else:
                p0.text = text
            changed.append(n)
        prs.save(sys.argv[4])
        print(f"노트를 바꾼 장 {len(changed)}: {changed}")
