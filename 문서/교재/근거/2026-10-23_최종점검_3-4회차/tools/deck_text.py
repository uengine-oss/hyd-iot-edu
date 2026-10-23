"""덱 PPTX의 장마다 화면 글(글상자 · 표 · 그룹 안 글)을 뽑는다. 발표자 노트는 뽑지 않는다."""
import sys
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
def walk(shapes):
    for sh in shapes:
        if sh.shape_type == MSO_SHAPE_TYPE.GROUP: yield from walk(sh.shapes)
        else: yield sh
def lines(sh):
    if sh.has_text_frame:
        t = " / ".join(p.text.strip() for p in sh.text_frame.paragraphs if p.text.strip())
        if t: yield t
    if getattr(sh, "has_table", False) and sh.has_table:
        for r in sh.table.rows:
            yield "| " + " | ".join(c.text.replace("\n", " / ").strip() for c in r.cells) + " |"
p = Presentation(sys.argv[1])
for i, s in enumerate(p.slides, 1):
    print(f"=== SLIDE {i}")
    for sh in walk(s.shapes):
        for t in lines(sh): print("  " + t)
