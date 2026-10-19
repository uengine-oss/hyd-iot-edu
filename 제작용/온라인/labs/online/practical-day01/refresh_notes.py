"""공통 화면·그림은 보존하고 발표자 노트만 최신 대본으로 바꾼다."""
from pathlib import Path
import re, zipfile, io, json, hashlib
from pptx import Presentation

ROOT=Path(__file__).resolve().parents[4]
deck=ROOT/'만들기/decks/practical/day01/실전반_01일차.md'
path=ROOT/'산출물/교재/PPTX/실전반_온라인/실전반_01일차.pptx'
before=path.read_bytes()
prs=Presentation(io.BytesIO(before))
blocks=re.split(r'(?=^## SLIDE )',deck.read_text('utf-8'),flags=re.M)[1:]
assert len(blocks)==len(prs.slides)==78
for block,slide in zip(blocks,prs.slides):
    match=re.search(r'^- 노트: (.*)$',block,re.M)
    slide.notes_slide.notes_text_frame.text=match[1] if match else ''
buffer=io.BytesIO();prs.save(buffer)
after=buffer.getvalue()
# python-pptx 저장 결과에서 노트 XML만 가져와 기존 ZIP의 나머지 항목을 그대로 둔다.
out=io.BytesIO();changed=[]
with zipfile.ZipFile(io.BytesIO(before)) as old,zipfile.ZipFile(io.BytesIO(after)) as new,zipfile.ZipFile(out,'w') as target:
    for item in old.infolist():
        data=old.read(item.filename)
        if re.fullmatch(r'ppt/notesSlides/notesSlide\d+\.xml',item.filename):
            replacement=new.read(item.filename)
            if replacement!=data:changed.append(item.filename)
            data=replacement
        target.writestr(item,data)
final=out.getvalue()
with zipfile.ZipFile(io.BytesIO(before)) as old,zipfile.ZipFile(io.BytesIO(final)) as new:
    assert new.testzip() is None
    for name in old.namelist():
        if not re.fullmatch(r'ppt/notesSlides/notesSlide\d+\.xml',name):assert old.read(name)==new.read(name),name
check=Presentation(io.BytesIO(final));norm=lambda s:re.sub(r'\s+','',s)
for block,slide in zip(blocks,check.slides):
    match=re.search(r'^- 노트: (.*)$',block,re.M)
    assert norm(slide.notes_slide.notes_text_frame.text)==norm(match[1] if match else '')
path.write_bytes(final)
(ROOT/'.evidence/practical01/notes-refresh-check.json').write_text(json.dumps({'slides':78,'notes_match':True,'changed_note_parts':len(changed),'all_other_zip_parts_unchanged':True,'prior_slide_render_still_applicable':True},indent=2),encoding='utf-8')
print('78 slides: notes match; every non-note ZIP part preserved byte-for-byte')
