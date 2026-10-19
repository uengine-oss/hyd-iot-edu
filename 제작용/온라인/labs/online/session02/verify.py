"""정적 문서 검증. 실습 코드 실행 없음."""
from pathlib import Path
import sys,re,json,hashlib,zipfile,posixpath,ast,importlib.util,difflib
from lxml import etree as E
from pptx import Presentation
from pypdf import PdfReader
import pypdfium2 as pdfium
from PIL import Image,ImageDraw,ImageChops
D=Path(__file__).resolve().parents[4];C=Path(sys.argv[1]);EV=D/'.evidence/session02'
I=C/'산출물/교재/대본/통합반_온라인/02회차';P=D/'산출물/교재/대본/실전반_온라인/02회차';B=EV/'before/산출물/교재/대본/통합반_온라인/02회차'
def write(p,s):p.parent.mkdir(parents=True,exist_ok=True);p.write_text(s,encoding='utf8')
def estimate(t):
 t=re.sub(r'^#.*$|^>.*$|\*\*','',t,flags=re.M);t=re.sub(r'[A-Za-z][A-Za-z0-9 _\-/\.]*\(([가-힣 ]+)\)',r'\1',t)
 return (len(re.findall('[가-힣]',t))+1.6*len(re.findall(r'\d',t))+.6*len(re.findall('[A-Za-z]',t)))/275
def speech(b):return ' '.join(l.strip().replace('**','') for l in b.splitlines()[1:] if l.strip() and not l.startswith(('>','#','---')))
def blocks(f):return re.split(r'(?=^## \[SLIDE )',f.read_text('utf8'),flags=re.M)[1:]
def theory(folder):return [b for f in sorted(folder.glob('*_이론_*.md')) for b in blocks(f)]
def notes(folder):
 out={}
 for b in theory(folder):
  m=re.match(r'## \[SLIDE (\d+)(?:~(\d+))?\]',b);a=int(m[1]);z=int(m[2] or m[1])
  for n in range(a,z+1):out[n]=speech(b) if n==a else f'{a}~{z}장 통합 발화는 {a}장 발표자 노트를 따른다.'
 return out
old=theory(B);now=theory(I);changed=[]
for a,b in zip(old,now):
 if speech(a)!=speech(b):changed.append(dict(block=b.splitlines()[0],before_minutes=estimate(speech(a)),after_minutes=estimate(speech(b))))
timing=dict(integrated=dict(slides_before=72,slides_after=72,blocks_before=len(old),blocks_after=len(now),changed_speech_blocks=len(changed),before_minutes=sum(estimate(speech(b)) for b in old),after_minutes=sum(estimate(speech(b)) for b in now),original_syllabus_theory_minutes=115,wait_seconds=145,practice=[dict(file=f.name,spoken_minutes=estimate(f.read_text('utf8'))) for f in sorted(I.glob('*_실습_*.md'))]),practical=dict(slides=107,blocks=len(theory(P)),minutes=sum(estimate(speech(b)) for b in theory(P)),original_syllabus_theory_minutes=95,wait_seconds=173,practice=[dict(file=f.name,spoken_minutes=estimate(f.read_text('utf8'))) for f in sorted(P.glob('*_실습_*.md'))]),integrated_changed_blocks=changed)
# 실전반의 공정한 비교 기준: 현재 통합반 72장 + 3회차 본문 4~33장 원발화. 구판 압축본과 비교하지 않는다.
source3=[]
for b in theory(C/'산출물/교재/대본/통합반_온라인/03회차'):
 n=int(re.match(r'## \[SLIDE (\d+)',b)[1])
 if 4<=n<=33:source3.append(b)
timing['practical']['full_source_minutes']=sum(estimate(speech(b)) for b in now+source3)
timing['practical']['source3_content_minutes']=sum(estimate(speech(b)) for b in source3)
timing['practical']['source2_minutes']=timing['integrated']['after_minutes']
timing['practical']['added_vs_full_source_seconds']=(timing['practical']['minutes']-timing['practical']['full_source_minutes'])*60
timing['integrated']['delta_seconds']=(timing['integrated']['after_minutes']-timing['integrated']['before_minutes'])*60
write(EV/'timing.json',json.dumps(timing,ensure_ascii=False,indent=2))
report={'lab_execution':False,'capture':False,'pptx':[],'preservation':{},'source_mapping':{}}
NS={'a':'http://schemas.openxmlformats.org/drawingml/2006/main','p':'http://schemas.openxmlformats.org/presentationml/2006/main'}
for root,folder,course,count in [(C,I,'통합반',72),(D,P,'실전반',107)]:
 path=root/f'산출물/교재/PPTX/{course}_온라인/{course}_02일차.pptx';nt=notes(folder);prs=Presentation(path);assert len(prs.slides)==count
 with zipfile.ZipFile(path) as z:
  assert z.testzip() is None
  for name in z.namelist():
   if name.endswith('.rels'):
    parent=posixpath.dirname(posixpath.dirname(name))
    for r in E.fromstring(z.read(name)):
     if r.get('TargetMode')=='External':continue
     target=r.get('Target');resolved=posixpath.normpath(posixpath.join(parent,target)) if not target.startswith('/') else target[1:];assert resolved in z.namelist(),(name,resolved)
  for n in range(1,count+1):
   rel=E.fromstring(z.read(f'ppt/slides/_rels/slide{n}.xml.rels'));target=next(r.get('Target') for r in rel if r.get('Type').endswith('/notesSlide'));nr=E.fromstring(z.read(posixpath.normpath(posixpath.join('ppt/slides',target))))
   body=nr.xpath('//p:sp[p:nvSpPr/p:nvPr/p:ph[@type="body"]]//a:t/text()',namespaces=NS);assert ''.join(body)==nt[n],(course,n,'notes mismatch')
 report['pptx'].append(dict(course=course,slides=count,notes_exact=True,zip_relationships_valid=True))
 # 学習者向け案内の参照を実在ファイルと照合。
 labs=sorted((root/f'산출물/실습/{course}_온라인').glob('02일차_실습*'))
 labs=[f for f in labs if f.is_dir()];assert len(labs)==(4 if course=='통합반' else 5)
 lines=[f'# {course} 2회차 실습 안내','', '자료를 다운로드해 압축을 풀고 안내.html부터 연다. 강사가 직접 제작·확인하는 흐름이며 실제 실행·캡처는 이번 제작에 포함하지 않았다. 접속 정보와 랩 DB는 수업 담당자가 별도로 제공한다.','']
 for lab in labs:
  with zipfile.ZipFile(lab.with_suffix('.zip')) as z:
   for f in lab.rglob('*'):
    if f.is_file():assert z.read(f.relative_to(lab).as_posix())==f.read_bytes()
  lines.append(f'- [{lab.name} 안내]({lab.name}/안내.html) · [다운로드]({lab.name}.zip) · [강사 대본]({lab.name}/강사대본.pdf)')
 write(root/f'산출물/실습/{course}_온라인/02회차_실습_안내.md','\n'.join(lines)+'\n')
# 통합반 화면 XML과 모든 그림·마스터·테마는 원본 바이트 그대로.
before=EV/'before/산출물/교재/PPTX/통합반_온라인/통합반_02일차.pptx';after=C/'산출물/교재/PPTX/통합반_온라인/통합반_02일차.pptx'
with zipfile.ZipFile(before) as a,zipfile.ZipFile(after) as b:
 same=[n for n in a.namelist() if re.fullmatch(r'ppt/slides/slide\d+.xml',n) or n.startswith(('ppt/media/','ppt/theme/','ppt/slideLayouts/','ppt/slideMasters/'))]
 assert all(a.read(n)==b.read(n) for n in same);report['preservation']['integrated_visible_parts_byte_identical']=len(same)
def deckblocks(path):return re.split(r'(?=^## SLIDE )',path.read_text('utf8'),flags=re.M)[1:]
src={d:deckblocks(C/f'만들기/decks/integrated/day{d:02}/통합반_{d:02}일차.md') for d in [2,3]};dst=deckblocks(D/'만들기/decks/practical/day02/실전반_02일차.md')
mapping=json.loads((P/'원본대응.json').read_text('utf8'))['slides'];verified=[]
def content(b):
 b=b.split('### 화면 구성')[0];b=re.sub(r'^## SLIDE \d+','## SLIDE',b);b=re.sub(r'^- 노트:.*$','',b,flags=re.M);b=re.sub(r'^- 번호:.*$','',b,flags=re.M);b=re.sub(r'figures/day\d+/', 'figures/',b);return b.strip()
for m in mapping:
 d=m['source_day'];n=m['source_slide'];dest=m['destination']
 if d and dest>3:assert content(src[d][n-1])==content(dst[dest-1]),(d,n,dest,'screen content changed');verified.append(dest)
assert sorted(m['source_slide'] for m in mapping if m['source_day']==2)==list(range(1,73))
assert sorted(m['source_slide'] for m in mapping if m['source_day']==3)==list(range(4,34))
report['source_mapping']=dict(source2_all72_once=True,source3_all30_content_once=True,common_content_exact=len(verified),metadata_adapted=3,new_slides=5)
# 공통 실습 발화의 과정명 외 차이가 없어야 한다.
for a,b in zip(sorted(I.glob('*_실습_*.md')),sorted(P.glob('*_실습_*.md'))[:4]):assert a.read_text('utf8').replace('통합반 2회차','실전반 2회차')==b.read_text('utf8')
for f in Path(__file__).parent.glob('*.py'):ast.parse(f.read_text('utf8'))
report['static_python_syntax']=True
# 전쪽 렌더·통합반 이전 화면 픽셀 비교. PDF 변환 완료 후에만 수행.
render_dir=EV/'final-render';renders=[]
if '--render' in sys.argv:
 for course,count in [('통합반',72),('실전반',107)]:
  f=render_dir/f'{course}_02일차.pdf';doc=pdfium.PdfDocument(f);assert len(doc)==count;folder=render_dir/course;folder.mkdir(exist_ok=True);thumbs=[]
  olddoc=pdfium.PdfDocument(EV/'source-render/통합반_02일차.pdf') if course=='통합반' else None
  for i in range(count):
   im=doc[i].render(scale=1).to_pil().convert('RGB');im.save(folder/f'{i+1:03}.png')
   if olddoc:
    oldim=olddoc[i].render(scale=1).to_pil().convert('RGB');assert ImageChops.difference(im,oldim).getbbox() is None,('visible difference',i+1)
   im.thumbnail((390,220));thumbs.append(im)
  for start in range(0,count,16):
   sheet=Image.new('RGB',(1600,1000),'#dce1e6');draw=ImageDraw.Draw(sheet)
   for j,im in enumerate(thumbs[start:start+16]):
    x=5+j%4*400;y=22+j//4*250;sheet.paste(im,(x,y));draw.text((x,y-17),str(start+j+1),fill='black')
   sheet.save(folder/f'contact-{start//16+1:02}.png')
  renders.append(dict(course=course,pages=count,integrated_all72_pixel_identical=bool(olddoc)))
 report['renders']=renders
baseline=json.loads((EV/'before.json').read_text('utf8'));changedpaths=[];missing=[]
for p,h in baseline.items():
 f=Path(p)
 if not f.exists():missing.append(p)
 elif hashlib.sha256(f.read_bytes()).hexdigest()!=h:changedpaths.append(p)
assert not missing,missing
allowed=['산출물/교재/대본/통합반_온라인/02회차/','산출물/교재/대본/통합반_온라인/통합반_02일차','산출물/교재/대본/통합반_온라인/낭독본/통합반_02일차','만들기/decks/integrated/day02/','산출물/교재/PPTX/통합반_온라인/통합반_02일차','문서/교재/인계_통합반_회차별_점검.md','문서/교재/4_만드는_법.md','문서/교재/3_현황.md','문서/교재/6_남은_일.md']
unexpected=[p for p in changedpaths if not any(s in p.replace('\\','/') for s in allowed)];assert not unexpected,unexpected
report['preservation'].update(baseline_files=len(baseline),changed_existing_files=len(changedpaths),missing_files=0,unexpected_changes=0,changed_paths=changedpaths)
write(EV/'final-check.json',json.dumps(report,ensure_ascii=False,indent=2))
print(json.dumps({'timing':timing,'checks':{k:v for k,v in report.items() if k!='preservation'}},ensure_ascii=False))
