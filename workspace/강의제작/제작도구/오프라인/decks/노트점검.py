"""오프라인 이론 덱의 발표자 노트 꼴을 점검한다(덱.md D-114 · D-144 · D-145). 형식 검사이고 윤문 판정이 아니다.

python3 workspace/강의제작/제작도구/오프라인/decks/노트점검.py <덱.md> <덱.pptx>
1 장마다 노트가 있고 [짚을 곳]으로 시작해 [다음]으로 끝난다
2 [짚을 곳] · [다음] 줄 말고는 모두 「·」로 시작한다(메모 안 「 / 」로 쪼개진 줄 없음)
3 PPTX 노트 = 덱 MD 노트(` / ` → 줄바꿈), 장 수 같음
4 첫 장 · 마지막 장 밖에서 실습 · 시연 예고 낱말이 없다(덱.md D-145)
"""
import re, sys
from pptx import Presentation
md, pp = sys.argv[1], sys.argv[2]
notes = [re.search(r'^- 노트: (.*)$', b, re.M) for b in re.split(r'(?=^## SLIDE )', open(md, encoding='utf8').read(), flags=re.M)[1:]]
errs = []
for i, m in enumerate(notes, 1):
    if not m: errs.append(f'{i}장 노트 없음'); continue
    lines = m.group(1).split(' / ')
    if not lines[0].startswith('[짚을 곳]'): errs.append(f'{i}장 [짚을 곳]으로 시작하지 않음')
    if not lines[-1].startswith('[다음]'): errs.append(f'{i}장 [다음]으로 끝나지 않음')
    for l in lines[1:-1]:
        if not l.startswith('· '): errs.append(f'{i}장 「·」 없는 줄: {l[:40]}')
    if 1 < i < len(notes) and re.search(r'실습에서|이어지는 실습|시연에서|다음 회차|핸즈온', m.group(1)):
        errs.append(f'{i}장 사이 장에 실습 · 회차 예고')
prs = Presentation(pp)
if len(prs.slides) != len(notes): errs.append(f'장 수 다름: MD {len(notes)} · PPTX {len(prs.slides)}')
for i, (s, m) in enumerate(zip(prs.slides, notes), 1):
    got = s.notes_slide.notes_text_frame.text.replace('\r', '\n').strip() if s.has_notes_slide else ''
    want = m.group(1).replace(' / ', '\n').strip() if m else ''
    if got != want: errs.append(f'{i}장 PPTX 노트가 덱 MD 노트와 다름')
print('\n'.join(errs) or f'노트 점검 문제 0 ({len(notes)}장)')
sys.exit(1 if errs else 0)
