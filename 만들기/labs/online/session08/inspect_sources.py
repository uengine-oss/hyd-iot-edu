from pathlib import Path
import json,re,hashlib
from openpyxl import load_workbook
D=Path(__file__).resolve().parents[4]; C=Path('D:/work/study/hyd-integrated-recovered-20261018'); E=D/'.evidence/session08'
rows=[]
for label,root in [('integrated',C),('practical',D)]:
 p=root/'문서/교재/실라버스/피지컬AI_실라버스_최종본.xlsx'; w=load_workbook(p,data_only=True)
 print(label,hashlib.sha256(p.read_bytes()).hexdigest(),w.sheetnames)
 for s in w:
  if '온라인' not in s.title:continue
  day=None
  for n,row in enumerate(s.values,1):
   vals=[str(x) if x is not None else '' for x in row]
   for x in vals[:3]:
    m=re.match(r'^(\d+)\s*(?:일|회)',x)
    if m:day=int(m[1])
   rows.append(dict(root=label,sheet=s.title,row=n,day=day,values=vals))
   if label=='practical' and (day in [8,11] or n<7):print(s.title,n,vals)
(E/'syllabus.json').write_text(json.dumps(rows,ensure_ascii=False,indent=2),encoding='utf8')
for day in [8,11]:
 p=C/f'만들기/decks/integrated/day{day:02}/통합반_{day:02}일차.md'
 t=re.sub(r'^- 노트:.*$','',p.read_text('utf8'),flags=re.M)
 (E/f'deck{day:02}-review.txt').write_text(t,encoding='utf8')
print('source decks extracted without duplicate notes')
